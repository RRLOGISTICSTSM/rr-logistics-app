import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
import os
import io
import re
import tempfile
import sqlite3

# PDF & Image Reading / Parsing
try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

from PIL import Image
try:
    import pytesseract
    HAS_PYTESSERACT = True
except ImportError:
    HAS_PYTESSERACT = False

from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# Safe import for openpyxl
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

st.set_page_config(page_title="R.R. Logistics TSM", page_icon="🚚", layout="wide")

# ==========================================
# 0. DATABASE & STORAGE SETUP
# ==========================================
UPLOAD_FOLDER = "saved_documents"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def init_db():
    conn = sqlite3.connect("rr_logistics.db")
    c = conn.cursor()
    
    # ॲप सेटिंग्ज
    c.execute('''CREATE TABLE IF NOT EXISTS app_settings (
                    key_name TEXT PRIMARY KEY,
                    key_value TEXT
                )''')
    c.execute("INSERT OR IGNORE INTO app_settings VALUES ('admin_password', 'rr123')")
    c.execute("INSERT OR IGNORE INTO app_settings VALUES ('security_pin', '9822')")

    # वाहनांचे टेबल
    c.execute('''CREATE TABLE IF NOT EXISTS vehicles (
                    veh_no TEXT PRIMARY KEY,
                    owner_name TEXT,
                    veh_type TEXT,
                    capacity TEXT,
                    default_dl TEXT,
                    rc_doc TEXT,
                    ins_doc TEXT,
                    ins_exp TEXT,
                    fit_doc TEXT,
                    fit_exp TEXT,
                    puc_doc TEXT,
                    puc_exp TEXT,
                    per_doc TEXT,
                    per_exp TEXT,
                    created_at TEXT
                )''')

    # ड्रायव्हर्स टेबल
    c.execute('''CREATE TABLE IF NOT EXISTS drivers (
                    dl_no TEXT PRIMARY KEY,
                    driver_name TEXT,
                    mobile TEXT,
                    aadhaar_no TEXT,
                    upi_id TEXT,
                    bank_name TEXT,
                    acc_no TEXT,
                    ifsc TEXT,
                    aadhaar_doc TEXT,
                    lic_doc TEXT,
                    created_at TEXT
                )''')

    # ट्रिप शीट्स टेबल
    c.execute('''CREATE TABLE IF NOT EXISTS trip_sheets (
                    lr_no TEXT PRIMARY KEY,
                    loading_date TEXT,
                    date_val TEXT,
                    from_loc TEXT,
                    to_loc TEXT,
                    truck_no TEXT,
                    driver_dl TEXT,
                    driver_name TEXT,
                    driver_mob TEXT,
                    rate_ton REAL,
                    weight_ton REAL,
                    freight REAL,
                    bill_no TEXT,
                    otp TEXT,
                    party_name TEXT,
                    material_grade TEXT,
                    start_km INTEGER,
                    end_km INTEGER,
                    running_km INTEGER,
                    trip_adv REAL,
                    unloading REAL,
                    food REAL,
                    tyre_bill REAL,
                    fastag1 REAL,
                    fastag2 REAL,
                    diesel_ltr REAL,
                    diesel_amt REAL,
                    commission REAL,
                    maintenance REAL,
                    other_exp REAL,
                    total_exp REAL,
                    trip_balance REAL,
                    created_at TEXT
                )''')

    # Schema Migration
    c.execute("PRAGMA table_info(trip_sheets)")
    existing_cols = [col[1] for col in c.fetchall()]
    for col_name, col_type in [("di_no", "TEXT DEFAULT ''"), ("bolt_seal_no", "TEXT DEFAULT ''"), ("material_grade", "TEXT DEFAULT 'PPC'")]:
        if col_name not in existing_cols:
            try:
                c.execute(f"ALTER TABLE trip_sheets ADD COLUMN {col_name} {col_type}")
            except:
                pass

    conn.commit()
    conn.close()

init_db()

def get_db_connection():
    return sqlite3.connect("rr_logistics.db")

def get_setting(key, default_val=""):
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT key_value FROM app_settings WHERE key_name=?", (key,))
    row = c.fetchone()
    conn.close()
    return row[0] if row else default_val

def set_setting(key, val):
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("INSERT OR REPLACE INTO app_settings VALUES (?, ?)", (key, str(val)))
    conn.commit()
    conn.close()

def fetch_vehicle_list():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT veh_no FROM vehicles ORDER BY veh_no", conn)
    conn.close()
    return df['veh_no'].tolist()

def fetch_driver_list():
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT dl_no, driver_name FROM drivers ORDER BY driver_name")
    rows = c.fetchall()
    conn.close()
    return rows

def get_vehicle_data(veh_no):
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT * FROM vehicles WHERE veh_no=?", (veh_no,))
    row = c.fetchone()
    conn.close()
    return row

def get_driver_data(dl_no):
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT * FROM drivers WHERE dl_no=?", (dl_no,))
    row = c.fetchone()
    conn.close()
    return row

def save_uploaded_file(uploaded_file, prefix, existing_path=""):
    if uploaded_file is None:
        return existing_path
    file_ext = os.path.splitext(uploaded_file.name)[1]
    safe_name = f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{file_ext}"
    file_path = os.path.join(UPLOAD_FOLDER, safe_name)
    with open(file_path, "wb") as f:
        f.write(uploaded_file.getbuffer())
    return file_path

# ==========================================
# PARSER: ULTRATECH DOCUMENT EXTRACTOR
# ==========================================
def extract_text_from_doc(file_obj):
    extracted_text = ""
    file_name = file_obj.name.lower()
    
    if file_name.endswith('.pdf'):
        if HAS_PDFPLUMBER:
            with pdfplumber.open(file_obj) as pdf:
                for page in pdf.pages:
                    txt = page.extract_text()
                    if txt:
                        extracted_text += txt + "\n"
        else:
            st.error("कृपया 'pdfplumber' इन्स्टॉल करा: pip install pdfplumber")
    elif file_name.endswith(('.jpg', '.jpeg', '.png')):
        if HAS_PYTESSERACT:
            image = Image.open(file_obj)
            extracted_text = pytesseract.image_to_string(image)
        else:
            st.warning("इमेज वाचण्यासाठी Tesseract उपलब्ध नाही. कृपया थेट PDF अपलोड करा किंवा Tesseract कॉन्फिगर करा.")
            
    return extracted_text

def parse_ultratech_content(text):
    data = {}
    if not text:
        return data

    # 1. DI No. / Dispatch Instruction
    di_match = re.search(r'(?:DI\s*(?:NO|Number)?[:\s\.-]+|DISPATCH\s*NO[:\s\.-]+)([A-Za-z0-9\/-]+)', text, re.IGNORECASE)
    if di_match:
        data['di_no'] = di_match.group(1).strip()

    # 2. Vehicle No. / Truck No.
    truck_match = re.search(r'([A-Z]{2}[-\s]?[0-9]{1,2}[-\s]?[A-Z]{1,3}[-\s]?[0-9]{4})', text)
    if truck_match:
        clean_truck = truck_match.group(1).replace(" ", "").replace("-", "")
        data['truck_no'] = clean_truck

    # 3. Bolt / Seal No.
    seal_match = re.search(r'(?:SEAL|BOLT|LOCK)\s*(?:NO)?[:\s\.-]+([A-Za-z0-9]+)', text, re.IGNORECASE)
    if seal_match:
        data['bolt_no'] = seal_match.group(1).strip()

    # 4. Weight in MT / Quantity
    weight_match = re.search(r'(?:NET\s*WT|WEIGHT|QTY|QUANTITY)[\s:\.-]+([0-9]+\.?[0-9]*)\s*(?:MT|TONS|TON)?', text, re.IGNORECASE)
    if weight_match:
        try:
            data['weight'] = str(float(weight_match.group(1)))
        except:
            pass

    # 5. E-Way Bill No.
    eway_match = re.search(r'(?:E[\s-]?WAY|EWB)\s*(?:BILL)?\s*(?:NO)?[:\s\.-]+([0-9]{12})', text, re.IGNORECASE)
    if eway_match:
        data['c_eway'] = eway_match.group(1).strip()

    # 6. GSTIN
    gst_match = re.search(r'([0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1})', text)
    if gst_match:
        data['c_gst'] = gst_match.group(1).strip()

    # 7. Material / Grade
    for grade in ["PPC", "OPC 43", "OPC 53", "SUPER", "COMPOSITE", "PSC"]:
        if re.search(rf'\b{grade}\b', text, re.IGNORECASE):
            data['mat_grade'] = grade
            break

    # 8. Date
    date_match = re.search(r'([0-3]?[0-9][\/\.-][0-1]?[0-9][\/\.-]20[2-9][0-9])', text)
    if date_match:
        data['date_val'] = date_match.group(1).replace('/', '-')

    # 9. Destination / To
    to_match = re.search(r'(?:DESTINATION|TO|CONSIGNEE\s*CITY)[:\s\.-]+([A-Za-z]+)', text, re.IGNORECASE)
    if to_match:
        data['to_loc'] = to_match.group(1).strip().upper()

    return data

# ==========================================
# 1. LR PDF GENERATION FUNCTION
# ==========================================
def generate_official_lr(lr_no, di_no, bolt_no, truck_no, date_val, from_loc, to_loc, consignor, consignee, c_gst, c_eway, d_name, d_dl, d_mobile, weight, freight, order_by, mat_grade):
    temp_dir = tempfile.gettempdir()
    safe_lr = str(lr_no).replace('/', '_')
    pdf_path = os.path.join(temp_dir, f"LR_{safe_lr}.pdf")
    
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    story = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=15, textColor=colors.HexColor('#1e3a8a'), alignment=1, spaceAfter=2)
    sub_style = ParagraphStyle('SubStyle', parent=styles['Normal'], fontSize=8.5, textColor=colors.HexColor('#1e3a8a'), alignment=1, spaceAfter=2)
    addr_style = ParagraphStyle('AddrStyle', parent=styles['Normal'], fontSize=7.5, textColor=colors.HexColor('#334155'), alignment=1, spaceAfter=8)
    
    for l_file in ["logo.jpeg", "logo.jpg", "logo.png"]:
        if os.path.exists(l_file):
            img = RLImage(l_file, width=80, height=55)
            img.hAlign = 'CENTER'
            story.append(img)
            story.append(Spacer(1, 4))
            break
    
    story.append(Paragraph("SUBJECT TO DHULE JURISDICTION", ParagraphStyle('SubTop', parent=styles['Normal'], fontSize=6.5, textColor=colors.HexColor('#475569'), alignment=1)))
    story.append(Paragraph("R.R. LOGISTICS TRANSPORT PVT. LTD.", title_style))
    story.append(Paragraph("TRANSPORT CONTRACTOR & COMMISSION AGENT", sub_style))
    story.append(Paragraph("Gat No. 298-2A, R.R. Petroleum, Mumbai-Agra Highway, Sarwad, Dhule.<br/><b>Mob: 9822414040 / 9049521177</b>", addr_style))
    
    data = [
        [
            Paragraph(f"<b>L.R. No. :</b> {lr_no}", styles['Normal']),
            Paragraph(f"<b>DI No. :</b> {di_no}<br/><b>Bolt / Seal No. :</b> {bolt_no}", styles['Normal']),
            Paragraph(f"<b>Date :</b> {date_val}<br/><b>Truck No. :</b> {truck_no}", styles['Normal'])
        ],
        [
            Paragraph(f"<b>From :</b> {from_loc}", styles['Normal']),
            "",
            Paragraph(f"<b>To :</b> {to_loc}", styles['Normal'])
        ],
        [
            Paragraph(f"<b>Consignor:</b><br/>{consignor}<br/><br/><b>GSTIN No. :</b> {c_gst}<br/><b>E-way Bill No. :</b> {c_eway}", styles['Normal']),
            "",
            Paragraph(f"<b>Consignee:</b><br/>{consignee}<br/><br/><b>Material / Grade:</b> {mat_grade}", styles['Normal'])
        ],
        [
            Paragraph(f"<b>Driver Name :</b> {d_name}", styles['Normal']),
            Paragraph(f"<b>D.L. No. :</b> {d_dl}", styles['Normal']),
            Paragraph(f"<b>Weight (MT) :</b> {weight}", styles['Normal'])
        ],
        [
            Paragraph(f"<b>Driver Mobile :</b> {d_mobile}", styles['Normal']),
            "",
            Paragraph(f"<b>Freight (भाडे) :</b> {freight}<br/><b>Order By :</b> {order_by}", styles['Normal'])
        ],
        [
            Paragraph("We are only Broker and commission agent. Please check all documents of truck carefully. Subject to DHULE Jurisdiction.", styles['Normal']),
            "",
            Paragraph("For R.R. LOGISTICS TRANSPORT PVT. LTD.<br/><br/><br/><b>Authorized Signatory</b>", styles['Normal'])
        ]
    ]
    
    t = Table(data, colWidths=[200, 185, 185])
    t.setStyle(TableStyle([
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#1e3a8a')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#1e3a8a')),
        ('SPAN', (0, 1), (1, 1)),
        ('SPAN', (0, 2), (1, 2)),
        ('SPAN', (0, 4), (1, 4)),
        ('SPAN', (0, 5), (1, 5)),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
    ]))
    
    story.append(t)
    story.append(Spacer(1, 4))
    story.append(Paragraph("THIS IS COMPUTER GENERATED BILTY", ParagraphStyle('Foot', parent=styles['Normal'], fontSize=6.5, textColor=colors.HexColor('#475569'), alignment=1)))
    
    doc.build(story)
    return pdf_path

# ==========================================
# 2. TRIP SHEET EXCEL (WITH FORMULAS)
# ==========================================
def generate_trip_sheet_excel(d):
    if not HAS_OPENPYXL:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trip Sheet"
    
    navy_fill = PatternFill(start_color="103A70", end_color="103A70", fill_type="solid")
    blue_header_fill = PatternFill(start_color="95B3D7", end_color="95B3D7", fill_type="solid")
    dark_blue_block = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "R. R. LOGISTICS, DHULE - TRIP SHEET"
    title_cell.font = font_title
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    headers_meta = [
        ("A2", "LR.NO", "B2", d['lr_no'], "C2", "DI NO", "D2", d.get('di_no', ''), "E2", "Date", "F2", d['date_val']),
        ("A3", "From", "B3", d['from_loc'], "C3", "To", "D3", d['to_loc'], "E3", "Truck No", "F3", d['truck_no']),
        ("A4", "Rate (per Ton)", "B4", d['rate_ton'], "C4", "Freight", "D4", "=B4*F4", "E4", "Weight (Tons)", "F4", d['weight_ton']),
        ("A5", "Bolt/Seal No", "B5", d.get('bolt_seal_no', ''), "C5", "Grade", "D5", d.get('material_grade', 'PPC'), "E5", "Party Name", "F5", d['party_name'])
    ]
    
    for row in headers_meta:
        ws[row[0]] = row[1]; ws[row[0]].font = font_bold
        ws[row[2]] = row[3]; ws[row[2]].font = font_regular
        ws[row[4]] = row[5]; ws[row[4]].font = font_bold
        ws[row[6]] = row[7]; ws[row[6]].font = font_regular
        ws[row[8]] = row[9]; ws[row[8]].font = font_bold
        ws[row[10]] = row[11]; ws[row[10]].font = font_regular
        
    ws.merge_cells("G2:H5")
    for r in range(2, 6):
        for c in range(7, 9):
            ws.cell(row=r, column=c).fill = dark_blue_block

    exp_headers = [("A7", "EXP"), ("B7", "Ltr"), ("C7", "Rate"), ("D7", "Amount"), ("E7", "Date"), ("F7", "Pump"), ("G7", "Summary"), ("H7", "Values")]
    for cell, val in exp_headers:
        ws[cell] = val
        ws[cell].fill = blue_header_fill
        ws[cell].font = font_bold
        ws[cell].alignment = Alignment(horizontal="center")

    exp_data = [
        ("Trip Adv", "", "", d['trip_adv'], d['adv_date'], ""),
        ("Unloading", "", "", d['unloading'], "", ""),
        ("Food", "", "", d['food'], "", ""),
        ("Tyre Bill", "", "", d['tyre_bill'], "", ""),
        ("FasTag-1", "", "", d['fastag1'], "", ""),
        ("FasTag-2", "", "", d['fastag2'], "", ""),
        ("Diesel-1", d['d1_ltr'], d['d1_rate'], "=B14*C14", d['d1_date'], d['d1_pump']),
        ("Diesel-2", d['d2_ltr'], d['d2_rate'], "=B15*C15", d['d2_date'], d['d2_pump']),
        ("Commission", "", "", d['commission'], "", ""),
        ("Maintenance", "", "", d['maintenance'], "", ""),
        ("Other", "", "", d['other_exp'], "", "")
    ]

    for idx, r_data in enumerate(exp_data, start=8):
        ws[f"A{idx}"] = r_data[0]
        ws[f"B{idx}"] = r_data[1]
        ws[f"C{idx}"] = r_data[2]
        ws[f"D{idx}"] = r_data[3]
        ws[f"E{idx}"] = r_data[4]
        ws[f"F{idx}"] = r_data[5]

    ws["A19"] = "Total Exp"
    ws["A19"].font = font_bold
    ws["D19"] = "=SUM(D8:D18)"
    ws["D19"].font = font_bold

    summary_data = [
        ("Start Km", d['start_km']),
        ("End Km", d['end_km']),
        ("Running Km", "=H9-H8"),
        ("Total Diesel (Ltr)", "=B14+B15"),
        ("Average (Km/L)", "=IF(H11>0, H10/H11, 0)"),
        ("Trip Balance", "=D4-D19"),
    ]
    for idx, (lbl, val) in enumerate(summary_data, start=8):
        ws[f"G{idx}"] = lbl
        ws[f"H{idx}"] = val
        ws[f"G{idx}"].font = font_bold

    ws["G19"] = "Trip Balance"
    ws["G19"].font = font_bold
    ws["H19"] = "=D4-D19"
    ws["H19"].font = font_bold

    for r in list(range(8, 14)) + list(range(16, 19)):
        ws[f"B{r}"].fill = dark_blue_block
        ws[f"C{r}"].fill = dark_blue_block

    for row in ws.iter_rows(min_row=1, max_row=19, min_col=1, max_col=8):
        for cell in row:
            cell.border = thin_border
            if cell.row in [2,3,4,5,7]:
                if not cell.alignment.horizontal:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            if str(cell.coordinate)[0] in ['B', 'C', 'D', 'H'] and cell.row >= 8:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

# ==========================================
# 3. TRIP SHEET PDF GENERATION
# ==========================================
def generate_trip_sheet_pdf(d):
    temp_dir = tempfile.gettempdir()
    safe_lr = str(d['lr_no']).replace('/', '_')
    pdf_path = os.path.join(temp_dir, f"TripSheet_{safe_lr}.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=landscape(letter), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    story = []
    
    freight_val = round(d['rate_ton'] * d['weight_ton'], 2)
    d1_amt = round(d['d1_ltr'] * d['d1_rate'], 2)
    d2_amt = round(d['d2_ltr'] * d['d2_rate'], 2)
    tot_exp = round(d['trip_adv'] + d['unloading'] + d['food'] + d['tyre_bill'] + d['fastag1'] + d['fastag2'] + d1_amt + d2_amt + d['commission'] + d['maintenance'] + d['other_exp'], 2)
    running_km = d['end_km'] - d['start_km']
    total_diesel = round(d['d1_ltr'] + d['d2_ltr'], 2)
    avg_kml = round(running_km / total_diesel, 2) if total_diesel > 0 else 0
    trip_bal = round(freight_val - tot_exp, 2)

    table_data = [
        ["R. R. LOGISTICS, DHULE - TRIP SHEET", "", "", "", "", "", "", ""],
        ["LR.NO", str(d['lr_no']), "DI NO", str(d.get('di_no', '')), "Date", str(d['date_val']), "", ""],
        ["From", str(d['from_loc']), "To", str(d['to_loc']), "Truck No", str(d['truck_no']), "", ""],
        ["Rate (per Ton)", str(d['rate_ton']), "Freight", str(freight_val), "Weight (Tons)", str(d['weight_ton']), "", ""],
        ["Bolt/Seal No", str(d.get('bolt_seal_no', '')), "Grade", str(d.get('material_grade', 'PPC')), "Party Name", str(d['party_name']), "", ""],
        ["EXP", "Ltr", "Rate", "Amount", "Date", "Pump", "Summary", "Values"],
        ["Trip Adv", "", "", str(d['trip_adv']), str(d['adv_date']), "", "Start Km", str(d['start_km'])],
        ["Unloading", "", "", str(d['unloading']), "", "", "End Km", str(d['end_km'])],
        ["Food", "", "", str(d['food']), "", "", "Running Km", str(running_km)],
        ["Tyre Bill", "", "", str(d['tyre_bill']), "", "", "Total Diesel (Ltr)", str(total_diesel)],
        ["FasTag-1", "", "", str(d['fastag1']), "", "", "Average (Km/L)", str(avg_kml)],
        ["FasTag-2", "", "", str(d['fastag2']), "", "", "Trip Balance", str(trip_bal)],
        ["Diesel-1", str(d['d1_ltr']), str(d['d1_rate']), str(d1_amt), str(d['d1_date']), str(d['d1_pump']), "", ""],
        ["Diesel-2", str(d['d2_ltr']), str(d['d2_rate']), str(d2_amt), str(d['d2_date']), str(d['d2_pump']), "", ""],
        ["Commission", "", "", str(d['commission']), "", "", "", ""],
        ["Maintenance", "", "", str(d['maintenance']), "", "", "", ""],
        ["Other", "", "", str(d['other_exp']), "", "", "", ""],
        ["Total Exp", "", "", str(tot_exp), "", "", "Trip Balance", str(trip_bal)]
    ]

    t = Table(table_data, colWidths=[90, 60, 60, 80, 85, 110, 110, 85])
    t.setStyle(TableStyle([
        ('SPAN', (0, 0), (7, 0)),
        ('BACKGROUND', (0, 0), (7, 0), colors.HexColor('#103A70')),
        ('TEXTCOLOR', (0, 0), (7, 0), colors.white),
        ('ALIGN', (0, 0), (7, 0), 'CENTER'),
        ('FONTNAME', (0, 0), (7, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (7, 0), 13),
        ('SPAN', (6, 1), (7, 4)),
        ('BACKGROUND', (6, 1), (7, 4), colors.HexColor('#4F81BD')),
        ('BACKGROUND', (0, 5), (7, 5), colors.HexColor('#95B3D7')),
        ('FONTNAME', (0, 5), (7, 5), 'Helvetica-Bold'),
        ('ALIGN', (0, 5), (7, 5), 'CENTER'),
        ('BACKGROUND', (1, 6), (2, 11), colors.HexColor('#4F81BD')),
        ('BACKGROUND', (1, 14), (2, 16), colors.HexColor('#4F81BD')),
        ('FONTNAME', (0, 1), (0, 4), 'Helvetica-Bold'),
        ('FONTNAME', (2, 1), (2, 4), 'Helvetica-Bold'),
        ('FONTNAME', (4, 1), (4, 4), 'Helvetica-Bold'),
        ('FONTNAME', (6, 6), (6, 11), 'Helvetica-Bold'),
        ('FONTNAME', (0, 17), (0, 17), 'Helvetica-Bold'),
        ('FONTNAME', (3, 17), (3, 17), 'Helvetica-Bold'),
        ('FONTNAME', (6, 17), (7, 17), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 6), (-1, -1), 'CENTER'),
    ]))
    
    story.append(t)
    doc.build(story)
    return pdf_path

# ==========================================
# 4. DOCUMENT EXPIRY PDF GENERATION
# ==========================================
def generate_doc_expiry_pdf(vehicle_rows):
    temp_dir = tempfile.gettempdir()
    pdf_path = os.path.join(temp_dir, f"Vehicle_Doc_Expiry_{datetime.now().strftime('%Y%m%d')}.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=landscape(letter), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    story = []
    styles = getSampleStyleSheet()
    
    title = Paragraph("<b>R.R. LOGISTICS TRANSPORT PVT. LTD.</b><br/><font size=10>गाड्यांच्या कागदपत्रांचा एक्सपायरी रिपोर्ट (Vehicle Documents Expiry Status)</font>", ParagraphStyle('ExpTitle', parent=styles['Heading1'], alignment=1, fontSize=14, textColor=colors.HexColor('#1e3a8a')))
    story.append(title)
    story.append(Spacer(1, 10))
    
    table_data = [
        ["गाडी क्रमांक", "मालकाचे नाव", "इन्शुरन्स एक्सपायरी", "फिटनेस एक्सपायरी", "PUC एक्सपायरी", "नॅशनल परमिट एक्सपायरी", "स्थिती (Status)"]
    ]
    
    today = datetime.now().date()
    for row in vehicle_rows:
        veh_no = row[0]
        owner = row[1]
        ins_exp = row[7] or "N/A"
        fit_exp = row[9] or "N/A"
        puc_exp = row[11] or "N/A"
        per_exp = row[13] or "N/A"
        
        alerts = []
        for exp_val, name in [(ins_exp, "Ins"), (fit_exp, "Fit"), (puc_exp, "PUC"), (per_exp, "Permit")]:
            if exp_val != "N/A":
                try:
                    exp_date = datetime.strptime(exp_val, "%Y-%m-%d").date()
                    if exp_date < today:
                        alerts.append(f"{name} Expired")
                    elif (exp_date - today).days <= 15:
                        alerts.append(f"{name} Due Soon")
                except:
                    pass
        status_text = ", ".join(alerts) if alerts else "सर्व वैध (Valid)"
        table_data.append([veh_no, owner, ins_exp, fit_exp, puc_exp, per_exp, status_text])
        
    t = Table(table_data, colWidths=[100, 120, 100, 100, 95, 120, 120])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#103A70')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t)
    doc.build(story)
    return pdf_path

# ==========================================
# 5. MAIN STREAMLIT APP
# ==========================================
def main():
    for l_file in ["logo.jpeg", "logo.jpg", "logo.png"]:
        if os.path.exists(l_file):
            col_l, col_m, col_r = st.columns([1, 4, 1])
            with col_m:
                st.image(l_file, width=120)
            break

    st.markdown("<h1 style='text-align: center; color: #1e3a8a;'>R.R. LOGISTICS TRANSPORT PVT. LTD.</h1>", unsafe_allow_html=True)
    st.markdown("<p style='text-align: center; color: #475569;'><b>TRANSPORT CONTRACTOR & COMMISSION AGENT</b> | Dhule Jurisdiction</p>", unsafe_allow_html=True)
    st.markdown("---")
    
    stored_password = get_setting("admin_password", "rr123")
    stored_security_pin = get_setting("security_pin", "9822")

    st.sidebar.title("🔐 सिस्टीम लॉगिन")
    username = st.sidebar.text_input("User ID", key="login_user")
    password = st.sidebar.text_input("Password", type="password", key="login_pass")
    
    if "logged_in" not in st.session_state:
        st.session_state.logged_in = False

    if st.sidebar.button("Login"):
        if username == "admin" and password == stored_password:
            st.session_state.logged_in = True
            st.sidebar.success("सक्सेसफुली लॉगिन झाले! ✅")
        else:
            st.sidebar.error("चुकीचा User ID किंवा Password!")

    with st.sidebar.expander("🔑 पासवर्ड बदला / विसरलात का?"):
        st.caption("पासवर्ड रीसेट करण्यासाठी सिक्युरिटी पिन टाका:")
        recovery_code = st.text_input("Security PIN", type="password", key="sec_pin")
        new_pass = st.text_input("नवीन पासवर्ड", type="password", key="new_pwd")
        confirm_pass = st.text_input("पासवर्ड खात्री करा", type="password", key="conf_pwd")
        
        if st.button("कायमस्वरूपी पासवर्ड अपडेट करा"):
            if recovery_code == stored_security_pin:
                if new_pass and new_pass == confirm_pass:
                    set_setting("admin_password", new_pass)
                    st.success("नवीन पासवर्ड सेव्ह झाला! ✅")
                else:
                    st.error("दोन्ही पासवर्ड जुळत नाहीत किंवा रिकामे आहेत!")
            else:
                st.error("चुकीचा Security PIN!")

    if st.session_state.logged_in:
        st.sidebar.markdown("---")
        menu = st.sidebar.radio("मेन्यू निवडा:", [
            "१. Vehicle Registration & Records", 
            "२. Driver Registration & Records", 
            "३. Trip Information & Auto LR Generation", 
            "४. Trip Sheet (Auto Excel & PDF)",
            "५. Saved Records & View Documents"
        ], key="dashboard_menu")
        
        existing_vehicles = fetch_vehicle_list()
        all_drivers = fetch_driver_list()
        driver_options = {d[0]: f"{d[1]} ({d[0]})" for d in all_drivers}

        # --- MENU 1: VEHICLE REGISTRATION ---
        if menu == "१. Vehicle Registration & Records":
            st.subheader("🚛 गाडी नोंदणी, माहिती संपादन व डॉक्युमेंट्स")
            
            selected_existing_veh = st.selectbox("नोंदणीकृत गाडी निवडा (माहिती पाहण्यासाठी/बदलण्यासाठी):", ["-- नवीन गाडी नोंदवा --"] + existing_vehicles)
            
            v_curr = None
            if selected_existing_veh != "-- नवीन गाडी नोंदवा --":
                v_curr = get_vehicle_data(selected_existing_veh)

            v_col1, v_col2 = st.columns(2)
            with v_col1:
                veh_no = st.text_input("गाडी क्रमांक (Vehicle Number)", value=v_curr[0] if v_curr else "MH18BZ2240")
                veh_owner = st.text_input("मालकाचे नाव (Owner Name)", value=v_curr[1] if v_curr else "")
            with v_col2:
                veh_type = st.text_input("गाडीचा प्रकार (Vehicle Type)", value=v_curr[2] if v_curr else "14 Wheeler")
                veh_cap = st.text_input("क्षमता टनांमध्ये (Capacity in Ton)", value=v_curr[3] if v_curr else "42")

            dl_keys = list(driver_options.keys())
            default_dl_idx = 0
            if v_curr and v_curr[4] in dl_keys:
                default_dl_idx = dl_keys.index(v_curr[4]) + 1

            sel_driver = st.selectbox("डिफॉल्ट ड्रायव्हर निवडा:", ["-- निवडा --"] + [driver_options[k] for k in dl_keys], index=default_dl_idx)
            sel_dl_code = ""
            if sel_driver != "-- निवडा --":
                sel_idx = (["-- निवडा --"] + [driver_options[k] for k in dl_keys]).index(sel_driver)
                sel_dl_code = dl_keys[sel_idx - 1]

            st.markdown("---")
            st.markdown("### 📂 कागदपत्रे आणि एक्सपायरी तारीख (Expiry Dates)")
            
            d1, d2 = st.columns(2)
            with d1:
                rc_file = st.file_uploader(f"1. RC Book {'(आधी सेव्ह आहे ✅)' if v_curr and v_curr[5] else ''}", type=["pdf", "jpg", "png"], key="rc_up")
                insurance_file = st.file_uploader(f"2. Insurance {'(आधी सेव्ह आहे ✅)' if v_curr and v_curr[6] else ''}", type=["pdf", "jpg", "png"], key="ins_up")
                ins_default_date = datetime.strptime(v_curr[7], "%Y-%m-%d").date() if v_curr and v_curr[7] else datetime.now().date() + timedelta(days=180)
                ins_exp = st.date_input("इन्शुरन्स एक्सपायरी", ins_default_date, key="ins_exp_key")
                
                fitness_file = st.file_uploader(f"3. Fitness Certificate {'(आधी सेव्ह आहे ✅)' if v_curr and v_curr[8] else ''}", type=["pdf", "jpg", "png"], key="fit_up")
                fit_default_date = datetime.strptime(v_curr[9], "%Y-%m-%d").date() if v_curr and v_curr[9] else datetime.now().date() + timedelta(days=180)
                fit_exp = st.date_input("फिटनेस एक्सपायरी", fit_default_date, key="fit_exp_key")
                
            with d2:
                puc_file = st.file_uploader(f"4. PUC {'(आधी सेव्ह आहे ✅)' if v_curr and v_curr[10] else ''}", type=["pdf", "jpg", "png"], key="puc_up")
                puc_default_date = datetime.strptime(v_curr[11], "%Y-%m-%d").date() if v_curr and v_curr[11] else datetime.now().date() + timedelta(days=90)
                puc_exp = st.date_input("PUC एक्सपायरी", puc_default_date, key="puc_exp_key")
                
                permit_file = st.file_uploader(f"5. National Permit {'(आधी सेव्ह आहे ✅)' if v_curr and v_curr[12] else ''}", type=["pdf", "jpg", "png"], key="per_up")
                per_default_date = datetime.strptime(v_curr[13], "%Y-%m-%d").date() if v_curr and v_curr[13] else datetime.now().date() + timedelta(days=365)
                per_exp = st.date_input("परमिट एक्सपायरी", per_default_date, key="per_exp_key")
            
            if st.button("गाडीची माहिती व कागदपत्रे कायम सेव्ह करा"):
                if veh_no:
                    rc_path = save_uploaded_file(rc_file, f"{veh_no}_RC", v_curr[5] if v_curr else "")
                    ins_path = save_uploaded_file(insurance_file, f"{veh_no}_INS", v_curr[6] if v_curr else "")
                    fit_path = save_uploaded_file(fitness_file, f"{veh_no}_FIT", v_curr[8] if v_curr else "")
                    puc_path = save_uploaded_file(puc_file, f"{veh_no}_PUC", v_curr[10] if v_curr else "")
                    per_path = save_uploaded_file(permit_file, f"{veh_no}_PER", v_curr[12] if v_curr else "")
                    
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute('''INSERT OR REPLACE INTO vehicles VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', 
                              (veh_no, veh_owner, veh_type, veh_cap, sel_dl_code, rc_path, ins_path, str(ins_exp), 
                               fit_path, str(fit_exp), puc_path, str(puc_exp), per_path, str(per_exp), datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
                    conn.commit()
                    conn.close()
                    st.success(f"गाडी क्र. {veh_no} ची माहिती कायमस्वरूपी सेव्ह झाली! ✅")
                    st.rerun()
                else:
                    st.warning("कृपया गाडी क्रमांक प्रविष्ट करा!")

        # --- MENU 2: DRIVER REGISTRATION ---
        elif menu == "२. Driver Registration & Records":
            st.subheader("👨‍✈️ ड्रायव्हर नोंदणी व बँक तपशील")
            
            sel_exist_dl = st.selectbox("नोंदणीकृत ड्रायव्हर निवडा (डेटा पाहण्यासाठी / बदलण्यासाठी):", ["-- नवीन ड्रायव्हर नोंदवा --"] + [f"{d[1]} - {d[0]}" for d in all_drivers])
            
            d_curr = None
            if sel_exist_dl != "-- नवीन ड्रायव्हर नोंदवा --":
                target_dl = sel_exist_dl.split(" - ")[-1]
                d_curr = get_driver_data(target_dl)

            d_col1, d_col2 = st.columns(2)
            with d_col1:
                drv_name = st.text_input("ड्रायव्हरचे नाव", value=d_curr[1] if d_curr else "Ramesh Pawar")
                drv_mob = st.text_input("मोबाईल नंबर", value=d_curr[2] if d_curr else "9823000000")
                drv_dl_no = st.text_input("ड्रायव्हिंग लायसन्स नंबर", value=d_curr[0] if d_curr else "MH1820120004567")
            with d_col2:
                drv_aad_no = st.text_input("आधार कार्ड नंबर", value=d_curr[3] if d_curr else "")
                drv_upi = st.text_input("UPI Number / GPay / PhonePe", value=d_curr[4] if d_curr else "")
            
            st.markdown("---")
            st.markdown("### 🏦 बँक तपशील आणि कागदपत्रे")
            b_col1, b_col2, b_col3 = st.columns(3)
            with b_col1:
                bank_name = st.text_input("बँकेचे नाव", value=d_curr[5] if d_curr else "")
            with b_col2:
                acc_no = st.text_input("खाते क्रमांक", value=d_curr[6] if d_curr else "")
            with b_col3:
                ifsc_code = st.text_input("IFSC कोड", value=d_curr[7] if d_curr else "")
            
            up_col1, up_col2 = st.columns(2)
            with up_col1:
                aadhar_file = st.file_uploader(f"Aadhaar Card Upload {'(आधी सेव्ह आहे ✅)' if d_curr and d_curr[8] else ''}", type=["pdf", "jpg", "png"], key="aadhar_up")
            with up_col2:
                licence_file = st.file_uploader(f"Driving Licence Upload {'(आधी सेव्ह आहे ✅)' if d_curr and d_curr[9] else ''}", type=["pdf", "jpg", "png"], key="lic_up")
            
            if st.button("ड्रायव्हरची संपूर्ण माहिती सेव्ह करा"):
                if drv_dl_no and drv_name:
                    aad_path = save_uploaded_file(aadhar_file, f"{drv_dl_no}_Aadhaar", d_curr[8] if d_curr else "")
                    lic_path = save_uploaded_file(licence_file, f"{drv_dl_no}_Licence", d_curr[9] if d_curr else "")
                    
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute('''INSERT OR REPLACE INTO drivers VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
                              (drv_dl_no, drv_name, drv_mob, drv_aad_no, drv_upi, bank_name, acc_no, ifsc_code, aad_path, lic_path, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
                    conn.commit()
                    conn.close()
                    st.success(f"ड्रायव्हर {drv_name} चा रेकॉर्ड डेटाबेसमध्ये कायम सेव्ह झाला! ✅")
                    st.rerun()
                else:
                    st.warning("कृपया ड्रायव्हरचे नाव व लायसन्स नंबर प्रविष्ट करा!")

        # --- MENU 3: TRIP INFO & AUTO LR GENERATION (WITH ULTRATECH SCANNER) ---
        elif menu == "३. Trip Information & Auto LR Generation":
            st.subheader("📄 ट्रिप माहिती आणि LR (Bilty) निर्मिती")

            # --- नवीन फिचर: अल्ट्राटेक डॉक्युमेंट अपलोड व ऑटो-स्कॅन ---
            st.markdown("### ⚡ अल्ट्राटेक सिस्टीम LR / इनव्हॉइस अपलोड करा (Auto Scan)")
            with st.expander("📥 अल्ट्राटेक PDF किंवा फोटो अपलोड करून फॉर्म आपोआप भरा", expanded=True):
                utcl_doc = st.file_uploader("Ultratech Bilty / Invoice ची PDF किंवा इमेज निवडा:", type=["pdf", "png", "jpg", "jpeg"], key="utcl_scanner")
                if utcl_doc is not None:
                    if st.button("🔍 डॉक्युमेंट स्कॅन करून फॉर्म भरा"):
                        with st.spinner("डॉक्युमेंटमधून माहिती काढली जात आहे..."):
                            extracted_txt = extract_text_from_doc(utcl_doc)
                            parsed_vals = parse_ultratech_content(extracted_txt)
                            if parsed_vals:
                                st.session_state['auto_utcl'] = parsed_vals
                                st.success("✅ डॉक्युमेंट यशस्वीरीत्या स्कॅन झाले! खालील फॉर्ममध्ये तपशील भरले गेले आहेत.")
                            else:
                                st.error("डॉक्युमेंटमधून माहिती वाचता आली नाही. कृपया स्पष्ट PDF किंवा फोटो अपलोड करा.")

            auto_data = st.session_state.get('auto_utcl', {})

            # गाडी निवडणे
            default_truck = auto_data.get('truck_no', '')
            truck_list = existing_vehicles if existing_vehicles else ["MH18BZ2240"]
            if default_truck and default_truck not in truck_list:
                truck_list = [default_truck] + truck_list

            sel_truck = st.selectbox("गाडी क्रमांक निवडा:", truck_list, key="lr_veh_sel")
            
            conn = get_db_connection()
            c = conn.cursor()
            c.execute("SELECT from_loc, to_loc, rate_ton, weight_ton, party_name, material_grade, di_no, bolt_seal_no FROM trip_sheets WHERE truck_no=? ORDER BY rowid DESC LIMIT 1", (sel_truck,))
            last_trip = c.fetchone()
            conn.close()
            
            v_data = get_vehicle_data(sel_truck)
            d_auto_name, d_auto_dl, d_auto_mob = "", "", ""
            if v_data and v_data[4]:
                d_info = get_driver_data(v_data[4])
                if d_info:
                    d_auto_dl = d_info[0]
                    d_auto_name = d_info[1]
                    d_auto_mob = d_info[2]

            col_lr1, col_lr2, col_lr3 = st.columns(3)
            with col_lr1:
                lr_no = st.text_input("L.R. No. (एल.आर. नंबर)", value="1602/1603")
            with col_lr2:
                default_di = auto_data.get('di_no', last_trip[6] if (last_trip and len(last_trip) > 6 and last_trip[6]) else "DI-99882")
                di_no = st.text_input("DI No. (डिस्पॅच सूचना क्र. / DI No.)", value=default_di)
            with col_lr3:
                default_bolt = auto_data.get('bolt_no', last_trip[7] if (last_trip and len(last_trip) > 7 and last_trip[7]) else "BT-4455")
                bolt_no = st.text_input("Bolt / Seal No. (बोल्ट सील क्रमांक)", value=default_bolt)

            col1, col2, col3 = st.columns(3)
            with col1:
                from_loc = st.text_input("From (कुठून)", value="ULTRATECH CEMENT LTD, DHULE")
                consignor = st.text_area("Consignor", value="Ultratech Cement Depo, Sarwad, Dhule")
                d_name = st.text_input("Driver Name", value=d_auto_name if d_auto_name else "Ramesh Pawar")
            with col2:
                truck_no = st.text_input("Truck No.", value=sel_truck)
                default_to = auto_data.get('to_loc', last_trip[1] if last_trip else "MALEGAON")
                to_loc = st.text_input("To (कुठे)", value=default_to)
                consignee = st.text_area("Consignee", value="Ultratech Authorised Dealer / Party Warehouse")
                d_dl = st.text_input("D.L. No.", value=d_auto_dl if d_auto_dl else "MH1820120004567")
            with col3:
                default_dt = auto_data.get('date_val', datetime.now().strftime("%d-%m-%Y"))
                date_val = st.text_input("Date", value=default_dt)
                default_gst = auto_data.get('c_gst', "27AABCU1234F1Z5")
                c_gst = st.text_input("Consignor GSTIN", value=default_gst)
                default_ewb = auto_data.get('c_eway', "EW-889923410")
                c_eway = st.text_input("E-way Bill No.", value=default_ewb)
                d_mobile = st.text_input("Driver Mobile No.", value=d_auto_mob if d_auto_mob else "9823000000")
            
            st.markdown("---")
            col_ex1, col_ex2, col_ex3, col_ex4 = st.columns(4)
            with col_ex1:
                default_wt = auto_data.get('weight', str(last_trip[3]) if last_trip else (v_data[3] if v_data and v_data[3] else "42.0"))
                weight = st.text_input("Weight (वजन टनांमध्ये)", value=default_wt)
            with col_ex2:
                grade_list = ["PPC", "Super", "OPC 43", "OPC 53", "PSC", "Composite", "Loose / Bulk"]
                detected_grade = auto_data.get('mat_grade', last_trip[5] if (last_trip and len(last_trip) > 5) else "PPC")
                default_grade_idx = 0
                for idx, g in enumerate(grade_list):
                    if g.upper() == detected_grade.upper():
                        default_grade_idx = idx
                        break
                mat_grade = st.selectbox("मटेरियल / ग्रेड निवडा (Material Grade):", grade_list, index=default_grade_idx)
            with col_ex3:
                rate_val = last_trip[2] if last_trip else 539.0
                try:
                    calc_f = float(weight) * float(rate_val)
                except:
                    calc_f = 0.0
                freight = st.text_input("एकूण भाडे (Freight ₹)", value=str(round(calc_f, 2)))
            with col_ex4:
                order_by = st.text_input("Order By", value=last_trip[4] if last_trip else "ULTRATECH")

            if st.button("🚀 LR PDF तयार करा"):
                if lr_no:
                    pdf_file_path = generate_official_lr(lr_no, di_no, bolt_no, truck_no, date_val, from_loc, to_loc, consignor, consignee, c_gst, c_eway, d_name, d_dl, d_mobile, weight, freight, order_by, mat_grade)
                    st.success("अधिकृत LR PDF तयार झाली आहे! 👇")
                    with open(pdf_file_path, "rb") as f:
                        st.download_button(
                            label="📥 डाऊनलोड ऑफिशियल LR PDF",
                            data=f,
                            file_name=f"LR_{lr_no.replace('/', '_')}_Official.pdf",
                            mime="application/pdf"
                        )
                else:
                    st.warning("कृपया L.R. No. प्रविष्ट करा!")

        # --- MENU 4: TRIP SHEET ---
        elif menu == "४. Trip Sheet (Auto Excel & PDF)":
            st.subheader("📊 R. R. LOGISTICS, DHULE - संपूर्ण ट्रिप शीट")
            st.caption("💡 डिझेल आणि सर्व ट्रिप खर्च थेट या शीटमध्ये नोंदवले जातात.")

            truck_choice = st.selectbox("गाडी क्रमांक निवडा:", existing_vehicles if existing_vehicles else ["MH18BZ2240"], key="ts_veh_sel")
            
            conn = get_db_connection()
            c = conn.cursor()
            c.execute("SELECT from_loc, to_loc, rate_ton, weight_ton, party_name, end_km, material_grade, di_no, bolt_seal_no FROM trip_sheets WHERE truck_no=? ORDER BY rowid DESC LIMIT 1", (truck_choice,))
            last_ts = c.fetchone()
            conn.close()

            col_ts_top1, col_ts_top2, col_ts_top3 = st.columns(3)
            with col_ts_top1:
                ts_lr = st.text_input("LR.NO", value="1602/1603")
            with col_ts_top2:
                ts_di = st.text_input("DI No.", value=last_ts[7] if (last_ts and len(last_ts) > 7 and last_ts[7]) else "DI-99882")
            with col_ts_top3:
                ts_bolt = st.text_input("Bolt / Seal No.", value=last_ts[8] if (last_ts and len(last_ts) > 8 and last_ts[8]) else "BT-4455")

            t1, t2, t3, t4 = st.columns(4)
            with t1:
                ts_from = st.text_input("From", value=last_ts[0] if last_ts else "ULTRATECH")
                ts_rate = st.number_input("Rate (per Ton)", min_value=0.0, value=float(last_ts[2]) if last_ts else 539.0, step=1.0)
                ts_bill = st.text_input("Bill No", value="")
            with t2:
                ts_ldate = st.text_input("Loading Date", value=datetime.now().strftime("%d-%m-%Y"))
                ts_to = st.text_input("To", value=last_ts[1] if last_ts else "MALEGAON")
                ts_weight = st.number_input("Weight (Tons)", min_value=0.0, value=float(last_ts[3]) if last_ts else 42.0, step=0.5)
            with t3:
                ts_date = st.text_input("Date", value=datetime.now().strftime("%d-%m-%Y"))
                ts_truck = st.text_input("Truck No", value=truck_choice)
                grade_options = ["PPC", "Super", "OPC 43", "OPC 53", "PSC", "Composite", "Loose / Bulk"]
                default_g_idx = 0
                if last_ts and len(last_ts) > 6 and last_ts[6] in grade_options:
                    default_g_idx = grade_options.index(last_ts[6])
                ts_grade = st.selectbox("Material / Grade", grade_options, index=default_g_idx)
            with t4:
                calc_freight = round(ts_rate * ts_weight, 2)
                st.metric("Freight (भाडे रक्कम)", f"₹ {calc_freight:,.2f}")
                ts_party = st.text_input("Party Name", value=last_ts[4] if last_ts else "DEPO")
                last_end_km = last_ts[5] if last_ts and last_ts[5] else 0
                ts_start_km = st.number_input("Start Km", min_value=0, value=last_end_km)
                ts_end_km = st.number_input("End Km", min_value=0, value=last_end_km + 250)
                calc_run_km = ts_end_km - ts_start_km
                st.metric("Running Km", f"{calc_run_km} KM")

            st.markdown("---")
            st.markdown("#### 💰 डिझेल व ट्रिप खर्च नोंद (Expenses & Diesel)")

            c_e1, c_e2, c_e3 = st.columns(3)
            with c_e1:
                exp_adv = st.number_input("Trip Adv (अ‍ॅडव्हान्स)", min_value=0.0, value=6000.0, step=500.0)
                adv_date = st.text_input("Adv Date", value=datetime.now().strftime("%d-%m-%Y"))
                exp_unloading = st.number_input("Unloading", min_value=0.0, value=0.0)
                exp_food = st.number_input("Food (जेवण)", min_value=0.0, value=0.0)
                exp_tyre = st.number_input("Tyre Bill", min_value=0.0, value=0.0)

            with c_e2:
                exp_ft1 = st.number_input("FasTag-1", min_value=0.0, value=0.0)
                exp_ft2 = st.number_input("FasTag-2", min_value=0.0, value=0.0)
                exp_comm = st.number_input("Commission", min_value=0.0, value=0.0)
                exp_maint = st.number_input("Maintenance", min_value=0.0, value=0.0)
                exp_other = st.number_input("Other Exp", min_value=0.0, value=0.0)

            with c_e3:
                st.markdown("**⛽ डिझेल तपशील (Diesel Entry)**")
                d1_ltr = st.number_input("Diesel-1 (Ltr)", min_value=0.0, value=90.0, step=5.0)
                d1_rate = st.number_input("Diesel-1 (Rate)", min_value=0.0, value=98.52, step=0.1)
                d1_amt = round(d1_ltr * d1_rate, 2)
                st.write(f"Diesel-1 Amount: ₹ {d1_amt}")
                d1_date = st.text_input("Diesel-1 Date", value=datetime.now().strftime("%d-%m-%Y"))
                d1_pump = st.text_input("Diesel-1 Pump", value="R.R. Petroleum, Sarwad")
                
                st.markdown("---")
                d2_ltr = st.number_input("Diesel-2 (Ltr)", min_value=0.0, value=0.0)
                d2_rate = st.number_input("Diesel-2 (Rate)", min_value=0.0, value=0.0)
                d2_amt = round(d2_ltr * d2_rate, 2)
                d2_date = st.text_input("Diesel-2 Date", value="")
                d2_pump = st.text_input("Diesel-2 Pump", value="")

            tot_diesel_ltr = round(d1_ltr + d2_ltr, 2)
            tot_exp = round(exp_adv + exp_unloading + exp_food + exp_tyre + exp_ft1 + exp_ft2 + d1_amt + d2_amt + exp_comm + exp_maint + exp_other, 2)
            trip_bal = round(calc_freight - tot_exp, 2)
            avg_km = round(calc_run_km / tot_diesel_ltr, 2) if tot_diesel_ltr > 0 else 0

            st.markdown("---")
            st.markdown("### 📋 ट्रिप सारांश (Auto Summary)")
            s1, s2, s3, s4 = st.columns(4)
            s1.metric("एकूण भाडे (Freight)", f"₹ {calc_freight:,.2f}")
            s2.metric("एकूण खर्च (Total Exp)", f"₹ {tot_exp:,.2f}")
            s3.metric("शिल्लक रक्कम (Trip Balance)", f"₹ {trip_bal:,.2f}")
            s4.metric("सरासरी (Average)", f"{avg_km} Km/L")

            ts_dict = {
                'lr_no': ts_lr, 'di_no': ts_di, 'bolt_seal_no': ts_bolt,
                'loading_date': ts_ldate, 'date_val': ts_date,
                'from_loc': ts_from, 'to_loc': ts_to, 'truck_no': ts_truck,
                'rate_ton': ts_rate, 'weight_ton': ts_weight,
                'bill_no': ts_bill, 'otp': "", 'party_name': ts_party,
                'material_grade': ts_grade,
                'start_km': ts_start_km, 'end_km': ts_end_km,
                'trip_adv': exp_adv, 'adv_date': adv_date,
                'unloading': exp_unloading, 'food': exp_food, 'tyre_bill': exp_tyre,
                'fastag1': exp_ft1, 'fastag2': exp_ft2,
                'd1_ltr': d1_ltr, 'd1_rate': d1_rate, 'd1_date': d1_date, 'd1_pump': d1_pump,
                'd2_ltr': d2_ltr, 'd2_rate': d2_rate, 'd2_date': d2_date, 'd2_pump': d2_pump,
                'commission': exp_comm, 'maintenance': exp_maint, 'other_exp': exp_other
            }

            if st.button("💾 ही ट्रिप शीट डेटाबेसमध्ये कायमस्वरूपी सेव्ह करा"):
                conn = get_db_connection()
                c = conn.cursor()
                c.execute('''INSERT OR REPLACE INTO trip_sheets (
                                lr_no, loading_date, date_val, from_loc, to_loc, truck_no, driver_dl, driver_name, driver_mob,
                                rate_ton, weight_ton, freight, bill_no, otp, party_name, material_grade, di_no, bolt_seal_no,
                                start_km, end_km, running_km, trip_adv, unloading, food, tyre_bill, fastag1, fastag2,
                                diesel_ltr, diesel_amt, commission, maintenance, other_exp, total_exp, trip_balance, created_at
                            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                          (ts_lr, ts_ldate, ts_date, ts_from, ts_to, ts_truck, "", "", "", ts_rate, ts_weight, calc_freight,
                           ts_bill, "", ts_party, ts_grade, ts_di, ts_bolt, ts_start_km, ts_end_km, calc_run_km, exp_adv, exp_unloading,
                           exp_food, exp_tyre, exp_ft1, exp_ft2, tot_diesel_ltr, (d1_amt+d2_amt), exp_comm, exp_maint,
                           exp_other, tot_exp, trip_bal, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
                conn.commit()
                conn.close()
                st.success(f"LR No. {ts_lr} ची ट्रिप शीट डेटाबेसमध्ये कायमस्वरूपी सेव्ह झाली! ✅")

            st.markdown("---")
            st.markdown("### 📥 फाईल डाऊनलोड करा (Excel & PDF)")
            down_col1, down_col2 = st.columns(2)
            
            with down_col1:
                if HAS_OPENPYXL:
                    excel_bytes = generate_trip_sheet_excel(ts_dict)
                    st.download_button(
                        label="📊 डाऊनलोड Trip Sheet Excel (with Live Formulas)",
                        data=excel_bytes,
                        file_name=f"TripSheet_{ts_lr.replace('/', '_')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                else:
                    df_ts = pd.DataFrame([ts_dict])
                    st.download_button(
                        label="📊 डाऊनलोड Trip Sheet CSV",
                        data=df_ts.to_csv(index=False).encode('utf-8'),
                        file_name=f"TripSheet_{ts_lr.replace('/', '_')}.csv",
                        mime="text/csv"
                    )

            with down_col2:
                pdf_file_path = generate_trip_sheet_pdf(ts_dict)
                with open(pdf_file_path, "rb") as f:
                    st.download_button(
                        label="📄 डाऊनलोड Trip Sheet PDF",
                        data=f,
                        file_name=f"TripSheet_{ts_lr.replace('/', '_')}.pdf",
                        mime="application/pdf"
                    )

        # --- MENU 5: SAVED RECORDS & EXPIRY REPORT ---
        elif menu == "५. Saved Records & View Documents":
            st.subheader("📁 सेव्ह केलेला डेटा, कागदपत्रे व डॉक्युमेंट एक्सपायरी रिपोर्ट")
            
            tab1, tab2, tab3 = st.tabs(["🚛 वाहने व डॉक्युमेंट एक्सपायरी", "📄 सेव्ह केलेल्या ट्रिप शीट्स", "👨‍✈️ ड्रायव्हर रेकॉर्ड्स"])
            conn = get_db_connection()
            
            with tab1:
                st.markdown("#### 🚚 वाहने व कागदपत्रे एक्सपायरी स्थिती")
                v_df = pd.read_sql_query("SELECT veh_no, owner_name, veh_type, ins_exp, fit_exp, puc_exp, per_exp FROM vehicles", conn)
                if not v_df.empty:
                    st.dataframe(v_df, use_container_width=True)
                    c = conn.cursor()
                    c.execute("SELECT * FROM vehicles")
                    v_all = c.fetchall()
                    exp_pdf = generate_doc_expiry_pdf(v_all)
                    with open(exp_pdf, "rb") as f:
                        st.download_button(
                            label="📥 डाऊनलोड सर्व गाड्यांचा एक्सपायरी PDF रिपोर्ट",
                            data=f,
                            file_name=f"Vehicle_Doc_Expiry_{datetime.now().strftime('%Y%m%d')}.pdf",
                            mime="application/pdf"
                        )
                else:
                    st.info("अद्याप कोणतीही गाडी नोंदणीकृत नाही.")
                    
                st.markdown("---")
                st.markdown("##### 🔍 गाडीचे सेव्ह कागदपत्र पहा / डाउनलोड करा:")
                sel_veh = st.selectbox("गाडी निवडा:", v_df['veh_no'].tolist() if not v_df.empty else [])
                if sel_veh:
                    c = conn.cursor()
                    c.execute("SELECT rc_doc, ins_doc, fit_doc, puc_doc, per_doc FROM vehicles WHERE veh_no=?", (sel_veh,))
                    docs = c.fetchone()
                    d_names = ["RC Book", "Insurance", "Fitness", "PUC", "National Permit"]
                    for name, path in zip(d_names, docs):
                        if path and os.path.exists(path):
                            with open(path, "rb") as file_data:
                                st.download_button(f"📥 डाउनलोड {name} ({os.path.basename(path)})", data=file_data, file_name=os.path.basename(path), key=f"dl_{name}_{sel_veh}")
                        else:
                            st.caption(f"{name}: कागदपत्र उपलब्ध नाही")

            with tab2:
                st.markdown("#### 📑 सेव्ह केलेल्या सर्व ट्रिप शीट्स")
                t_df = pd.read_sql_query("SELECT lr_no, di_no, bolt_seal_no, date_val, truck_no, from_loc, to_loc, material_grade, freight, total_exp, trip_balance, created_at FROM trip_sheets", conn)
                if not t_df.empty:
                    st.dataframe(t_df, use_container_width=True)
                else:
                    st.info("अद्याप कोणतीही ट्रिप शीट सेव्ह केलेली नाही.")

            with tab3:
                st.markdown("#### 👨‍✈️ ड्रायव्हर यादी व बँक तपशील")
                d_df = pd.read_sql_query("SELECT dl_no, driver_name, mobile, bank_name, acc_no, ifsc, created_at FROM drivers", conn)
                if not d_df.empty:
                    st.dataframe(d_df, use_container_width=True)
                else:
                    st.info("अद्याप कोणत्याही ड्रायव्हरची नोंदणी नाही.")
            
            conn.close()

    else:
        st.info("कृपया सिस्टीममध्ये प्रवेश करण्यासाठी युझर आयडी आणि पासवर्ड टाकून लॉगिन करा.")

if __name__ == '__main__':
    main()
